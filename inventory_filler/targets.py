"""Where the plan actually lands: a Google Sheet, or a local .xlsx.

Both targets expose the same three operations - read the grid, apply cell
edits, highlight cells that want a human eye - so the planner never learns
which one it is writing to, and a Google setup is never required to use the
tool.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from .planner import CellEdit, column_letter

REVIEW_COLOR = "FFF3B0"
"""Pale amber. Loud enough to find by eye, quiet enough to leave in the sheet."""

SHEETS_SCOPES = ["https://www.googleapis.com/auth/spreadsheets"]


def cell_value(edit: CellEdit) -> object:
    """Prices belong in the sheet as numbers so they still sum."""
    if edit.field_name == "price":
        try:
            return int(edit.value)
        except (TypeError, ValueError):
            return edit.value
    return edit.value


def quote_sheet_name(name: str) -> str:
    """A1 notation needs the tab name quoted, with internal quotes doubled."""
    return "'" + name.replace("'", "''") + "'"


def split_spec(spec: str) -> tuple[str, str]:
    """Split "book:in.xlsx#在庫" into its location and tab name."""
    location, _, sheet_name = spec.partition("#")
    return location, sheet_name


class WorkbookTarget:
    """A local .xlsx file, edited in place."""

    def __init__(self, path: str | Path, sheet_name: str = "", date_format: str = "%Y/%m/%d"):
        self.path = Path(path)
        self.sheet_name = sheet_name
        self.date_format = date_format
        self._workbook = None

    def describe(self) -> str:
        return f"{self.path}" + (f" [{self.sheet_name}]" if self.sheet_name else "")

    def _sheet(self):
        if self._workbook is None:
            try:
                from openpyxl import load_workbook
            except ImportError as exc:  # optional dependency
                raise ImportError("xlsx を扱うには `pip install openpyxl` が必要です") from exc

            self._workbook = load_workbook(self.path)
        if self.sheet_name:
            if self.sheet_name not in self._workbook.sheetnames:
                raise KeyError(
                    f"シート '{self.sheet_name}' が見つかりません。"
                    f"候補: {', '.join(self._workbook.sheetnames)}"
                )
            return self._workbook[self.sheet_name]
        return self._workbook.active

    def _text(self, value: object) -> str:
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.strftime(self.date_format)
        if isinstance(value, date):
            return value.strftime(self.date_format)
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    def read_grid(self) -> list[list[str]]:
        return [[self._text(cell.value) for cell in row] for row in self._sheet().iter_rows()]

    def apply(self, edits: list[CellEdit]) -> None:
        sheet = self._sheet()
        for edit in edits:
            sheet.cell(row=edit.row, column=edit.column).value = cell_value(edit)
        self._workbook.save(self.path)

    def highlight(self, cells: list[tuple[int, int]]) -> None:
        if not cells:
            return
        from openpyxl.styles import PatternFill

        fill = PatternFill(start_color=REVIEW_COLOR, end_color=REVIEW_COLOR, fill_type="solid")
        sheet = self._sheet()
        for row, column in cells:
            sheet.cell(row=row, column=column).fill = fill
        self._workbook.save(self.path)


class SheetsTarget:
    """A Google Sheet, reached with a service account."""

    def __init__(self, spreadsheet_id: str, sheet_name: str = "", credentials_path: str = ""):
        self.spreadsheet_id = spreadsheet_id
        self.sheet_name = sheet_name
        self.credentials_path = credentials_path
        self._service = None
        self._sheet_id: int | None = None

    def describe(self) -> str:
        return f"Googleスプレッドシート {self.spreadsheet_id}" + (
            f" [{self.sheet_name}]" if self.sheet_name else ""
        )

    def _api(self):
        if self._service is None:
            if not self.credentials_path:
                raise ValueError(
                    "サービスアカウントの鍵が指定されていません "
                    "(--credentials か GOOGLE_APPLICATION_CREDENTIALS)"
                )
            try:
                from google.oauth2 import service_account
                from googleapiclient.discovery import build
            except ImportError as exc:  # optional dependency
                raise ImportError(
                    "Googleスプレッドシートに書き込むには "
                    "`pip install google-api-python-client google-auth` が必要です"
                ) from exc

            credentials = service_account.Credentials.from_service_account_file(
                self.credentials_path, scopes=SHEETS_SCOPES
            )
            self._service = build("sheets", "v4", credentials=credentials)
        return self._service.spreadsheets()

    def _resolve_sheet(self) -> tuple[str, int]:
        """Look up the tab's title and numeric id, defaulting to the first tab."""
        metadata = self._api().get(spreadsheetId=self.spreadsheet_id).execute()
        tabs = metadata.get("sheets") or []
        if not tabs:
            raise RuntimeError("スプレッドシートにシートが1枚もありません")
        for tab in tabs:
            properties = tab.get("properties", {})
            if not self.sheet_name or properties.get("title") == self.sheet_name:
                return properties["title"], properties["sheetId"]
        titles = ", ".join(t.get("properties", {}).get("title", "?") for t in tabs)
        raise KeyError(f"シート '{self.sheet_name}' が見つかりません。候補: {titles}")

    def read_grid(self) -> list[list[str]]:
        title, self._sheet_id = self._resolve_sheet()
        self.sheet_name = title
        response = (
            self._api()
            .values()
            .get(spreadsheetId=self.spreadsheet_id, range=quote_sheet_name(title))
            .execute()
        )
        return [[str(cell).strip() for cell in row] for row in response.get("values", [])]

    def apply(self, edits: list[CellEdit]) -> None:
        if not edits:
            return
        title = self.sheet_name or self._resolve_sheet()[0]
        data = [
            {
                "range": f"{quote_sheet_name(title)}!{column_letter(edit.column)}{edit.row}",
                "values": [[cell_value(edit)]],
            }
            for edit in edits
        ]
        self._api().values().batchUpdate(
            spreadsheetId=self.spreadsheet_id,
            # USER_ENTERED so a date lands as a date and a price as a number,
            # exactly as if it had been typed in.
            body={"valueInputOption": "USER_ENTERED", "data": data},
        ).execute()

    def highlight(self, cells: list[tuple[int, int]]) -> None:
        if not cells:
            return
        if self._sheet_id is None:
            _, self._sheet_id = self._resolve_sheet()
        red, green, blue = (int(REVIEW_COLOR[i : i + 2], 16) / 255 for i in (0, 2, 4))
        requests = [
            {
                "repeatCell": {
                    "range": {
                        "sheetId": self._sheet_id,
                        "startRowIndex": row - 1,
                        "endRowIndex": row,
                        "startColumnIndex": column - 1,
                        "endColumnIndex": column,
                    },
                    "cell": {
                        "userEnteredFormat": {
                            "backgroundColor": {"red": red, "green": green, "blue": blue}
                        }
                    },
                    "fields": "userEnteredFormat.backgroundColor",
                }
            }
            for row, column in cells
        ]
        self._api().batchUpdate(
            spreadsheetId=self.spreadsheet_id, body={"requests": requests}
        ).execute()


def open_target(spec: str, *, credentials_path: str = "", date_format: str = "%Y/%m/%d"):
    """Build a target from "book:<path>[#tab]" or "sheet:<id>[#tab]"."""
    location, sheet_name = split_spec(spec)
    kind, _, rest = location.partition(":")
    if kind == "book":
        return WorkbookTarget(rest, sheet_name, date_format=date_format)
    if kind == "sheet":
        return SheetsTarget(rest, sheet_name, credentials_path=credentials_path)
    raise ValueError(
        f"書き込み先の指定が不正です: {spec!r} "
        "(book:<ファイル.xlsx> または sheet:<スプレッドシートID> の形式)"
    )
