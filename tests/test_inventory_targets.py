import json
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

from inventory_filler.cli import run
from inventory_filler.config import SheetConfig
from inventory_filler.planner import CellEdit
from inventory_filler.targets import (
    SheetsTarget,
    WorkbookTarget,
    cell_value,
    open_target,
    quote_sheet_name,
    split_spec,
)

HEADER = ["管理番号", "商品名", "仕入日", "仕入価格", "状態", "要確認"]


def _workbook(tmp_path, *rows, sheet_title="在庫"):
    path = tmp_path / "在庫表.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = sheet_title
    sheet.append(HEADER)
    for row in rows:
        sheet.append(list(row))
    book.save(path)
    return path


def _edit(row, column, value, field_name="price", confidence=1.0):
    return CellEdit(
        row=row,
        column=column,
        value=value,
        field_name=field_name,
        management_no="A-1",
        confidence=confidence,
    )


def test_cell_value_keeps_prices_numeric_so_the_column_still_sums():
    assert cell_value(_edit(2, 4, "1200")) == 1200
    assert cell_value(_edit(2, 3, "2024/03/05", field_name="purchased_on")) == "2024/03/05"


def test_cell_value_falls_back_to_text_for_an_unparsable_price():
    assert cell_value(_edit(2, 4, "応相談")) == "応相談"


def test_quote_sheet_name_doubles_internal_quotes():
    assert quote_sheet_name("在庫") == "'在庫'"
    assert quote_sheet_name("Bob's") == "'Bob''s'"


def test_split_spec_separates_the_tab_name():
    assert split_spec("book:在庫表.xlsx#在庫") == ("book:在庫表.xlsx", "在庫")
    assert split_spec("sheet:abc123") == ("sheet:abc123", "")


def test_open_target_builds_each_backend():
    assert isinstance(open_target("book:a.xlsx"), WorkbookTarget)
    assert isinstance(open_target("sheet:abc123#在庫"), SheetsTarget)


def test_open_target_rejects_an_unknown_scheme():
    with pytest.raises(ValueError, match="書き込み先の指定が不正"):
        open_target("https://docs.google.com/spreadsheets/d/abc")


def test_workbook_target_reads_the_grid_as_text(tmp_path):
    path = _workbook(tmp_path, ["A-1", "花瓶", datetime(2024, 3, 5), 1200.0, "A", ""])
    grid = WorkbookTarget(path, "在庫").read_grid()

    assert grid[0] == HEADER
    # Dates and whole floats must read back in the sheet's own formatting, or a
    # re-run would see a difference that isn't there and report a false conflict.
    assert grid[1][2] == "2024/03/05"
    assert grid[1][3] == "1200"


def test_workbook_target_applies_edits_and_saves(tmp_path):
    path = _workbook(tmp_path, ["A-1", "花瓶", "", "", "", ""])
    target = WorkbookTarget(path, "在庫")
    target.read_grid()
    target.apply([_edit(2, 4, "1200"), _edit(2, 5, "A", field_name="condition")])

    sheet = load_workbook(path)["在庫"]
    assert sheet.cell(row=2, column=4).value == 1200
    assert sheet.cell(row=2, column=5).value == "A"


def test_workbook_target_highlights_cells_for_review(tmp_path):
    path = _workbook(tmp_path, ["A-1", "花瓶", "", "", "", ""])
    target = WorkbookTarget(path, "在庫")
    target.read_grid()
    target.highlight([(2, 4)])

    sheet = load_workbook(path)["在庫"]
    assert sheet.cell(row=2, column=4).fill.fill_type == "solid"
    assert sheet.cell(row=2, column=3).fill.fill_type != "solid"


def test_workbook_target_names_the_tabs_it_does_have(tmp_path):
    path = _workbook(tmp_path, sheet_title="在庫")
    with pytest.raises(KeyError, match="在庫"):
        WorkbookTarget(path, "ない名前").read_grid()


def test_sheets_target_requires_credentials():
    with pytest.raises(ValueError, match="サービスアカウント"):
        SheetsTarget("abc123").read_grid()


def test_config_roundtrip(tmp_path):
    path = tmp_path / "inventory.json"
    SheetConfig(key_column="番号", confidence_threshold=0.6, overwrite=True).save(path)

    loaded = SheetConfig.load(path)
    assert loaded.key_column == "番号"
    assert loaded.confidence_threshold == 0.6
    assert loaded.overwrite is True


def test_config_load_missing_file_returns_defaults(tmp_path):
    config = SheetConfig.load(tmp_path / "nope.json")
    assert config.key_column == "管理番号"
    assert config.overwrite is False


def _saved_read(tmp_path, no="A-2", confidence=0.95):
    path = tmp_path / "read.json"
    path.write_text(
        json.dumps(
            [
                {
                    "source_image": "sticker.jpg",
                    "management_no": {"value": no, "confidence": confidence},
                    "purchased_on": {"value": "R6.3.5", "confidence": confidence},
                    "price": {"value": "￥1,200-", "confidence": confidence},
                    "condition": {"value": "美品", "confidence": confidence},
                }
            ],
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    return path


def test_cli_dry_run_changes_nothing(tmp_path, capsys):
    book = _workbook(tmp_path, ["A-2", "皿", "", "", "", ""])
    config = tmp_path / "inventory.json"
    SheetConfig(condition_aliases={"美品": "A"}).save(config)

    exit_code = run(
        ["--from-read", str(_saved_read(tmp_path)), "--target", f"book:{book}#在庫",
         "--config", str(config)]
    )

    assert exit_code == 0
    assert "--apply" in capsys.readouterr().out
    assert load_workbook(book)["在庫"].cell(row=2, column=4).value is None


def test_cli_apply_writes_the_normalized_values(tmp_path):
    book = _workbook(tmp_path, ["A-2", "皿", "", "", "", ""])
    config = tmp_path / "inventory.json"
    SheetConfig(condition_aliases={"美品": "A"}).save(config)

    exit_code = run(
        ["--from-read", str(_saved_read(tmp_path)), "--target", f"book:{book}#在庫",
         "--config", str(config), "--apply"]
    )

    assert exit_code == 0
    sheet = load_workbook(book)["在庫"]
    assert sheet.cell(row=2, column=3).value == "2024/03/05"
    assert sheet.cell(row=2, column=4).value == 1200
    assert sheet.cell(row=2, column=5).value == "A"


def test_cli_apply_flags_low_confidence_cells_in_the_sheet(tmp_path):
    book = _workbook(tmp_path, ["A-2", "皿", "", "", "", ""])
    config = tmp_path / "inventory.json"
    SheetConfig(review_column="要確認").save(config)

    run(["--from-read", str(_saved_read(tmp_path, confidence=0.4)),
         "--target", f"book:{book}#在庫", "--config", str(config), "--apply"])

    sheet = load_workbook(book)["在庫"]
    assert sheet.cell(row=2, column=6).value.startswith("要確認:")
    assert sheet.cell(row=2, column=4).fill.fill_type == "solid"


def test_cli_stops_when_a_configured_column_is_missing(tmp_path, capsys):
    book = _workbook(tmp_path, ["A-2", "皿", "", "", "", ""])
    config = tmp_path / "inventory.json"
    SheetConfig(columns={"price": "存在しない見出し"}).save(config)

    exit_code = run(
        ["--from-read", str(_saved_read(tmp_path)), "--target", f"book:{book}#在庫",
         "--config", str(config), "--apply"]
    )

    assert exit_code == 1
    assert "見つからない列" in capsys.readouterr().out
    assert load_workbook(book)["在庫"].cell(row=2, column=4).value is None


def test_cli_init_config_writes_a_starter_file(tmp_path, capsys):
    config = tmp_path / "inventory.json"
    assert run(["--init-config", "--config", str(config)]) == 0
    assert SheetConfig.load(config).key_column == "管理番号"
    assert run(["--init-config", "--config", str(config)]) == 1  # refuses to clobber
